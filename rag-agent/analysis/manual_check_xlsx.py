#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
"""Human-verification workbooks for the two corpora with no annotated header
tree. Generates only; nothing here scores anything."""
from __future__ import annotations

import json
import random
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "analysis"))

from header_path_coverage import TEMPLATE, load_corpus               # noqa: E402
from qwen_equiv_k import A, args_for                                 # noqa: E402
from rag_agent.serialization.caption import caption_sentence         # noqa: E402

SEED, N_GOLD, N_RAND = 42, 50, 50
PAD = 10                       # cells kept on each side of the target
XLSX_LIMIT = 32767
RENDER_DIR = Path("results/audit/renders")
FILL_COLS = ["값_맞음", "행경로_맞음", "열경로_맞음", "비고"]
COLS = ["no", "dataset", "table_id", "row", "col", "is_gold", "cell_sentence",
        "caption", "row_path", "col_path", "cell_value", "table_render"] + FILL_COLS


def render(C, tid, i, j, nhr_lines):
    """The table as text with row/col numbers, the target wrapped in [[ ]]."""
    lines = C.md_lines[tid]
    n_r, n_c = C.shape[tid]
    # md_lines = header block + separator + one line per DATA row
    first = len(lines) - n_r
    head = lines[:first]
    body = lines[first:]
    grid = [[p.strip() for p in ln.split("|")[1:-1]] for ln in body]
    width = max((len(r) for r in grid), default=0)
    off = width - n_c                       # leading header columns
    lo_r, hi_r = max(0, i - PAD), min(n_r, i + PAD + 1)
    lo_c, hi_c = max(0, j - PAD), min(n_c, j + PAD + 1)
    cut_r = (lo_r > 0) or (hi_r < n_r)
    cut_c = (lo_c > 0) or (hi_c < n_c)

    keep_c = list(range(off)) + [off + c for c in range(lo_c, hi_c)]
    out = ["# header rows (never cut):"]
    out += head
    if cut_c:
        out.append(f"...(중략: 열 {lo_c}~{hi_c - 1} 만 표시, 전체 {n_c} 열)")
    out.append("")
    out.append("col#   | " + " | ".join(
        ("h" + str(c) if c < off else str(c - off)) for c in keep_c))
    if lo_r > 0:
        out.append("...(중략)")
    for r in range(lo_r, hi_r):
        row = grid[r] if r < len(grid) else []
        cells = []
        for c in keep_c:
            v = row[c] if c < len(row) else ""
            if r == i and c == off + j:
                v = f"[[{v}]]"
            cells.append(v)
        out.append(f"row {r:<4} | " + " | ".join(cells))
    if hi_r < n_r:
        out.append("...(중략)")
    return "\n".join(out)


def build(dataset, population, out_path):
    if dataset == "multihiertt":
        a = A()
        a.dataset, a.population, a.rhb_question_types = dataset, population, []
    else:
        a = args_for(dataset, population)
    C = load_corpus(a)
    by = defaultdict(dict)
    for n, (t, i, j) in enumerate(C.cell_owner):
        by[t][(i, j)] = n

    gold = sorted({(str(t), int(i), int(j)) for q in C.queries
                   for (t, i, j) in q["gold_cells"]
                   if t in by and (i, j) in by[t]})
    allc = sorted({(t, i, j) for t, d in by.items() for (i, j) in d})
    rnd = random.Random(SEED)
    g_pick = rnd.sample(gold, min(N_GOLD, len(gold)))
    pool = [c for c in allc if c not in set(g_pick)]
    r_pick = rnd.sample(pool, min(N_RAND, len(pool)))

    RENDER_DIR.mkdir(parents=True, exist_ok=True)
    rows = []
    for no, (cell, is_gold) in enumerate(
            [(c, "Y") for c in g_pick] + [(c, "N") for c in r_pick], 1):
        tid, i, j = cell
        n = by[tid][(i, j)]
        rp, cp, v = C.cell_paths[n]
        sent = caption_sentence(C.title.get(tid, ""), rp, cp, v,
                                template=TEMPLATE["S3c"])
        tr = render(C, tid, i, j, None)
        if len(tr) > XLSX_LIMIT:
            f = RENDER_DIR / f"{dataset}_{tid}_{i}_{j}.txt"
            f.write_text(tr, encoding="utf-8")
            tr = str(f)
        rows.append([no, dataset, tid, i, j, is_gold, sent,
                     C.title.get(tid, ""), " > ".join(map(str, rp)),
                     " > ".join(map(str, cp)), str(v), tr] + [""] * len(FILL_COLS))

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "manual_check"
    ws.append(COLS)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for c in range(1, len(COLS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF2CC")
    first_fill = COLS.index(FILL_COLS[0]) + 1
    for c in range(first_fill, len(COLS) + 1):
        for r in range(1, len(rows) + 2):
            ws.cell(row=r, column=c).fill = fill
    widths = {"no": 5, "dataset": 13, "table_id": 22, "row": 6, "col": 6,
              "is_gold": 8, "cell_sentence": 70, "caption": 30, "row_path": 30,
              "col_path": 30, "cell_value": 14, "table_render": 90}
    for k, w in widths.items():
        ws.column_dimensions[chr(64 + COLS.index(k) + 1)].width = w
    for k in FILL_COLS:
        ws.column_dimensions[chr(64 + COLS.index(k) + 1)].width = 12
    wrap = Alignment(wrap_text=True, vertical="top")
    for k in ("cell_sentence", "table_render"):
        c = COLS.index(k) + 1
        for r in range(2, len(rows) + 2):
            ws.cell(row=r, column=c).alignment = wrap
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"file": str(out_path), "n_rows": len(rows),
            "n_gold": sum(1 for r in rows if r[5] == "Y"),
            "n_random": sum(1 for r in rows if r[5] == "N"),
            "gold_pool": len(gold), "cell_pool": len(allc),
            "n_render_offloaded": sum(1 for r in rows
                                      if str(r[11]).startswith("results/"))}


def main() -> int:
    out = [build("realhitbench", "rhb_lookup_all",
                 "results/audit/manual_check_rhb.xlsx"),
           build("multihiertt", "", "results/audit/manual_check_mh.xlsx")]
    print(json.dumps(out, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
