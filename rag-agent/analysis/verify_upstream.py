# SPDX-License-Identifier: MIT
"""Differential checks against frozen public source files, without loading LLMs.

Pure upstream corpus/rendering methods and isolated control flow are executed. This does not claim
to run the upstream retrieval models, SQL service, or complete QA systems.
"""
import argparse
import ast
import json
import sys
import typing
from collections import Counter
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import pandas as pd
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import Workbook, load_workbook
from rag_agent.eval.artifacts import digest, file_digest
from rag_agent.serialization.chunks import markdown_source
from scripts import retrieval_accuracy as ra


class Document:
    def __init__(self, page_content, metadata=None):
        self.page_content = page_content
        self.metadata = metadata or {}


def upstream_class(path, name, methods, namespace):
    namespace.update({"Tuple": tuple, "List": list})
    tree = ast.parse(path.read_text(encoding="utf-8"))
    node = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == name)
    node.body = [n for n in node.body if isinstance(n, ast.FunctionDef) and n.name in methods]
    module = ast.fix_missing_locations(ast.Module(body=[node], type_ignores=[]))
    exec(compile(module, str(path), "exec"), namespace)
    return namespace[name].__new__(namespace[name])


def verify(source_dir, huawei_dir, data_dir):
    source_dir, huawei_dir = Path(source_dir), Path(huawei_dir)
    manifest = json.loads((source_dir / "sources.json").read_text())
    for name, repo in manifest.items():
        for filename, sha in repo["files"].items():
            if file_digest(source_dir / name / filename) != sha:
                raise ValueError(f"changed upstream snapshot: {name}/{filename}")
    retriever = source_dir / "google/table_rag/agent/retriever.py"
    native = upstream_class(retriever, "Retriever", {
        "build_schema_corpus", "build_cell_corpus", "build_row_corpus", "build_column_corpus",
        "sample_rows_and_columns"}, {"Document": Document, "Counter": Counter})
    native.max_encode_cell = 10000
    frame = pd.DataFrame({"stub": ["r0", "r1"], "Revenue": [10, 30], "Cost": [20, 40]})
    native.df = frame
    native.row_retriever = SimpleNamespace(invoke=lambda q: [Document("", {"row_id": 0})])
    native.column_retriever = SimpleNamespace(invoke=lambda q: [Document("", {"col_id": 1})])
    subtable = native.sample_rows_and_columns("revenue?")
    assert list(subtable.columns) == ["Revenue"] and subtable.iloc[0, 0] == 10
    schemas = native.build_schema_corpus(frame)
    assert schemas[1].page_content == "Revenue" and "min" in schemas[1].metadata["result_text"]
    numeric_docs = [d.page_content for d in native.build_cell_corpus(frame)]
    assert any('"dtype": "int64"' in text for text in numeric_docs)
    class Table:
        data = [["10", "20"], ["30", "40"]]
        n_rows = n_cols = 2
        def row_path(self, i): return [f"r{i}"]
        def col_path(self, j): return [["Revenue"], ["Cost"]][j]
    selection = ra.rowcol_select([0, 1], [{("T", 0, 0), ("T", 0, 1)},
                                         {("T", 0, 0), ("T", 1, 0)}],
                                 ["row", "column"], [True, False], 1, 1,
                                 grid={"T": (Table(), "Title")})
    local_render = "\n".join(selection.context)
    assert all(str(c) in local_render for c in subtable.columns)
    assert "10" in local_render and "30" not in local_render and "20" not in local_render
    google = {"commit": manifest["google"]["commit"],
              "sampled_markdown": subtable.to_markdown(index=False),
              "t2_sampled_markdown": local_render,
              "selected_values_and_headers_match": True,
              "schema_index_text": [d.page_content for d in schemas],
              "schema_reader_text": [d.metadata["result_text"] for d in schemas],
              "cell_documents": numeric_docs,
              "column_index_text": [d.page_content for d in native.build_column_corpus(frame)]}
    huawei_file = huawei_dir / "online_inference/tools/retriever.py"
    chunker = upstream_class(huawei_file, "MixedDocRetriever", {"nltk_single_doc_chunking"}, {})
    chunker.text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    queries = ra.load_queries(data_dir, "test", {})
    tids = sorted({q["table_id"] for q in queries})
    mismatch, live, delivered = [], set(), set()
    for tid in tids:
        tab = ra.hg.load_table(tid, data_dir)
        source, _ = markdown_source(tab, tab.table, tab.title)
        expected, _ = chunker.nltk_single_doc_chunking(source, tab.title)
        actual = ra.trag_hetero_chunks(tab, tab.table, tab.title, 1000, 200)
        if [x[0] for x in actual] != expected:
            mismatch.append(tid)
        live |= {(tid, i, j) for i in range(tab.table.n_rows) for j in range(tab.table.n_cols)
                 if str(tab.table.data[i][j]).strip()}
        delivered |= {(tid, i, j) for _, cells in actual for i, j in cells}
    assert not mismatch
    assert delivered <= live
    # Execute the original Excel renderer with a small in-memory workbook.
    util = huawei_dir / "online_inference/utils/tool_utils.py"
    node = next(n for n in ast.parse(util.read_text()).body
                if isinstance(n, ast.FunctionDef) and n.name == "excel_to_markdown")
    workbook = Workbook()
    workbook.active.append(["stub", "Revenue", "Cost"])
    workbook.active.append(["r0", "10", None])
    env = {"load_workbook": lambda path: workbook}
    exec(compile(ast.Module(body=[node], type_ignores=[]), str(util), "exec"), env)
    native_render = env["excel_to_markdown"]("/fixture/table.xlsx")
    workbook.active.cell(2, 2).value = 10
    try:
        env["excel_to_markdown"]("/fixture/table.xlsx")
    except TypeError:
        numeric_excel_error = True
    else:
        numeric_excel_error = False
    # Exercise upstream default reranker initialization without loading weights.
    # All substitutes are model/device factories; control flow is upstream's.
    class Model:
        def eval(self): return self
        def to(self, device): return self
    factory = SimpleNamespace(from_pretrained=lambda *a, **kw: Model())
    fake_torch = SimpleNamespace(device=lambda name: name,
                                 no_grad=lambda: (lambda fn: fn),
                                 cuda=SimpleNamespace(is_available=lambda: True))
    reranker = upstream_class(util, "Reranker", {"__init__"},
                              {"Union": typing.Union, "torch": fake_torch,
                               "AutoTokenizer": factory, "AutoModelForSequenceClassification": factory})
    reranker.__init__(model_name_or_path="fixture")
    assert reranker.num_gpus == 1
    assert reranker.device == "cuda:4"
    # Run the exact SQL response assignment/exception handler in isolation.
    main_file = huawei_dir / "online_inference/main.py"
    source_tree = ast.parse(main_file.read_text())
    sql_try = next(n for n in ast.walk(source_tree) if isinstance(n, ast.Try)
                   and any(isinstance(x, ast.Name) and x.id == "get_excel_rag_response" for x in ast.walk(n)))
    try:
        exec(compile(ast.Module(body=[sql_try], type_ignores=[]), str(main_file), "exec"),
             {"get_excel_rag_response": lambda: None,
              "excel_rag_response": {"sql_str": "SELECT 1", "sql_execution_result": "1", "nl2sql_prompt": "schema"}})
    except ValueError as error:
        sql_handler_error = str(error)
    else:
        raise AssertionError("upstream SQL handler behavior changed")
    return {"sources": manifest, "google_pure_method_checks": google,
            "huawei": {"retriever_sha256": file_digest(huawei_file),
                       "native_splitter_tables": len(tids), "different_tables": mismatch,
                       "full_source_cells": len(live), "covered_cells": len(delivered),
                       "missing_cells": sorted(live - delivered),
                       "native_excel_render": native_render,
                       "native_numeric_excel_TypeError": numeric_excel_error,
                       "default_reranker_initialization": {"device": reranker.device,
                                                           "num_gpus": reranker.num_gpus},
                       "sql_response_handler_error": sql_handler_error,
                       "utils_sha256": file_digest(util), "main_sha256": file_digest(main_file)},
            "scope": "Pure upstream methods and local adapters; no full model/SQL/QA reproduction"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sources-dir", required=True)
    parser.add_argument("--huawei-dir", required=True)
    parser.add_argument("--data-dir", default="data/hitab")
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    result = verify(args.sources_dir, args.huawei_dir, args.data_dir)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("x", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(json.dumps({"splitter_tables": result["huawei"]["native_splitter_tables"],
                      "mismatches": len(result["huawei"]["different_tables"]),
                      "missing_cells": len(result["huawei"]["missing_cells"])}))
