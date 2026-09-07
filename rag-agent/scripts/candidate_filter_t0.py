"""T0 for PREREG-2026-08-28-candidate-filter.md — can a second signal find the
gold sentence among the 12-14 that are already in the context?

No LLM, no GPU. Reads a context workbook written by ``error_worksheet.py
--also-correct`` (whose replay refuses to write when it disagrees with the
recorded run) and re-scores each query's OWN context sentences with a
cross-encoder. Nothing is retrieved, reordered or refilled here: the question is
only whether the gold sentence would SURVIVE a keep-top-k cut.

The gate is retention, not accuracy -- §2 of the prereg: dropping the gold costs
~.72 EM on that query while purifying buys at most .13, so the filter is only
worth building if it almost never drops the gold.
"""
from __future__ import annotations

import argparse
import json
from collections import defaultdict
from pathlib import Path

import openpyxl


def rows_of(ws):
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def retained(scores, gold_idx, keep):
    """Indices of the top-``keep`` scores, and whether every gold one survives."""
    order = sorted(range(len(scores)), key=lambda i: -scores[i])[:keep]
    return order, gold_idx <= set(order)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--contexts", required=True, help="xlsx from error_worksheet.py --also-correct")
    ap.add_argument("--reranker", default="BAAI/bge-reranker-large")
    ap.add_argument("--keep", type=int, default=3)
    ap.add_argument("--batch", type=int, default=32)
    ap.add_argument("--max-length", type=int, default=256)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.contexts, read_only=True)
    meta = {r[0]: r[1] for r in wb["meta"].iter_rows(min_row=1, values_only=True)} \
        if "meta" in wb.sheetnames else {}
    qinfo = {r["query_id"]: r for r in rows_of(wb["errors"])}
    ctx = defaultdict(list)
    for r in rows_of(wb["context"]):
        ctx[r["query_id"]].append(r)

    # OSC=1 only: where the gold is absent there is nothing for a filter to keep.
    qids = [q for q in ctx if int(qinfo[q]["osc"] or 0) == 1]
    pairs, spans = [], []
    for q in qids:
        cells = ctx[q]
        spans.append((q, len(pairs), len(cells)))
        pairs += [(qinfo[q]["question"], c["sentence"]) for c in cells]
    print(f"[t0] {len(qids)} OSC=1 queries / {len(pairs)} pairs")

    from sentence_transformers import CrossEncoder
    ce = CrossEncoder(args.reranker, max_length=args.max_length)
    scores = ce.predict(pairs, batch_size=args.batch, show_progress_bar=True)

    n_top1 = n_keep = n_top1_bi = 0
    per_q = []
    for q, off, k in spans:
        cells = ctx[q]
        gold = {i for i, c in enumerate(cells) if int(c["is_gold"] or 0) == 1}
        if not gold:                       # replay says OSC=1, so this cannot happen
            raise SystemExit(f"{q}: OSC=1 but no gold sentence in its context")
        s = [float(x) for x in scores[off:off + k]]
        order, kept = retained(s, gold, args.keep)
        bi = [float(c["cos_q_sentence"]) for c in cells]
        n_top1 += order[0] in gold
        n_keep += kept
        n_top1_bi += max(range(k), key=lambda i: bi[i]) in gold
        per_q.append({"query_id": q, "n_ctx": k, "ce_top1_is_gold": order[0] in gold,
                      "gold_kept": bool(kept),
                      "gold_ce_rank": min(sorted(range(k), key=lambda i: -s[i]).index(g)
                                          for g in gold) + 1})
    n = len(spans)
    t0a, t0b, base = n_top1 / n, n_keep / n, n_top1_bi / n
    out = {
        "prereg": "PREREG-2026-08-28-candidate-filter.md §3",
        "contexts": args.contexts, "reranker": args.reranker, "keep": args.keep,
        "n_osc1": n, "n_pairs": len(pairs),
        "replay_drift": meta.get("replay_drift", meta.get("drift")),
        "T0-A_ce_top1_is_gold": round(t0a, 4),
        "T0-B_gold_kept_at_keep": round(t0b, 4),
        "T0-D_lift_over_bi_encoder": round(t0a - base, 4),
        "bi_encoder_top1_is_gold": round(base, 4),
        "gates": {"T0-A": t0a >= .75, "T0-B": t0b >= .95, "T0-D": (t0a - base) >= .10},
        "verdict": ("PROCEED to T1" if (t0b >= .95 and (t0a - base) >= .10) else
                    "RETRY at keep=5" if t0b >= .90 and (t0a - base) >= .10 else
                    "STOP -- close the context-composition axis"),
        "per_query": per_q,
    }
    Path(args.out).write_text(json.dumps(out, ensure_ascii=False, indent=1))
    print(json.dumps({k: v for k, v in out.items() if k != "per_query"},
                     ensure_ascii=False, indent=1))
    return 0


def _selftest():
    # gold survives a keep-2 cut only when its score is in the top 2
    assert retained([.9, .1, .8], {2}, 2) == ([0, 2], True)
    assert retained([.9, .1, .8], {1}, 2)[1] is False
    assert retained([.9, .1, .8], {0, 2}, 2)[1] is True     # both gold cells kept
    assert retained([.9, .1, .8], {0, 1}, 2)[1] is False    # one gold dropped -> OSC lost
    print("ok")


if __name__ == "__main__":
    import sys
    raise SystemExit(_selftest() if "--selftest" in sys.argv else main())
